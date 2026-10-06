"""Independent python-docx source for the inherited table appearance audit."""
import argparse
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def element(name, **values):
    item = OxmlElement("w:" + name)
    for key, value in values.items():
        item.set(qn("w:" + key), str(value))
    return item


parser = argparse.ArgumentParser()
parser.add_argument("output")
parser.add_argument("--conditional", action="store_true")
parser.add_argument("--text", action="store_true")
args = parser.parse_args()
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1
title = doc.styles["Title"]
title.font.name = "Arial"
title.font.size = Pt(20)
title.font.color.rgb = RGBColor(0, 0, 0)
title.font.underline = False
# The bundled template contains a decorative Title bottom border. It is not
# part of this table fixture; remove it rather than exporting that decoration.
for item in list(title.element.get_or_add_pPr()):
    if item.tag == qn("w:pBdr"):
        title.element.pPr.remove(item)
doc.add_paragraph("Table text formatting audit" if args.text else "Conditional table appearance audit" if args.conditional else "Inherited table appearance audit", "Title")
doc.add_paragraph("This report tests table style inheritance and explicit overrides during visual editing and save reopening.")

base = doc.styles.add_style("AuditTableBase", WD_STYLE_TYPE.TABLE)
properties = element("tblPr")
borders = element("tblBorders")
for side in ["top", "bottom", "left", "right", "insideH", "insideV"]:
    borders.append(element(side, val="single", sz=8, color="D9D9D9"))
properties.append(borders)
margins = element("tblCellMar")
for side, width in [("top", 60), ("bottom", 60), ("left", 160), ("right", 160)]:
    margins.append(element(side, w=width, type="dxa"))
properties.append(margins)
properties.append(element("tblW", w=7200, type="dxa"))
properties.append(element("tblLayout", type="fixed"))
base.element.append(properties)
child = doc.styles.add_style("AuditTableChild", WD_STYLE_TYPE.TABLE)
child.base_style = base
properties = element("tblPr")
margins = element("tblCellMar")
margins.append(element("left", w=80, type="dxa"))
properties.append(margins)
child.element.append(properties)
cell_defaults = element("tcPr")
cell_defaults.append(element("shd", val="clear", fill="EDF2F8"))
child.element.append(cell_defaults)

if args.text:
    # Keep an explicit non-Normal paragraph style in table cells. Native Word's
    # special Normal/docDefaults equivalence is a separate unresolved fixture,
    # not an excuse to guess or weaken this table-text acceptance case.
    doc.styles.add_style("AuditTableBody", WD_STYLE_TYPE.PARAGRAPH)
    run_properties = element("rPr")
    run_properties.append(element("rFonts", ascii="Arial", hAnsi="Arial"))
    run_properties.append(element("sz", val=22))
    run_properties.append(element("color", val="334455"))
    run_properties.append(element("b", val=0))
    base.element.append(run_properties)
    paragraph_properties = element("pPr")
    paragraph_properties.append(element("jc", val="left"))
    paragraph_properties.append(element("spacing", before=0, after=60, line=240, lineRule="auto"))
    base.element.append(paragraph_properties)
    for kind in ["firstRow", "lastCol"]:
        override = element("tblStylePr", type=kind)
        run_properties = element("rPr")
        run_properties.append(element("b", val=1 if kind == "firstRow" else 0))
        run_properties.append(element("color", val="000000" if kind == "firstRow" else "334455"))
        if kind == "firstRow":
            run_properties.append(element("sz", val=24))
        override.append(run_properties)
        paragraph_properties = element("pPr")
        paragraph_properties.append(element("jc", val="center"))
        override.append(paragraph_properties)
        # Reuse the existing region below instead of generating ambiguous roots.
        child.element.append(override)

if args.conditional:
    # These rules belong to real table styles, not direct cell shading. The
    # independent source therefore cannot inherit Fountain's importer mistakes.
    properties = base.element.find(qn("w:tblPr"))
    properties.append(element("tblStyleRowBandSize", val=1))
    properties.append(element("tblStyleColBandSize", val=1))
    for kind, fill in [
        ("band1Horz", "F5F8FD"), ("band2Horz", "FFFFFF"),
        ("firstCol", "EFEFEF"), ("lastCol", "E7E6E6"),
        ("firstRow", "CFE2F3"), ("lastRow", "DDEEDD"),
        ("nwCell", "C9DAF8"), ("neCell", "D9EAD3"),
        ("swCell", "FFF2CC"), ("seCell", "F4CCCC"),
    ]:
        override = element("tblStylePr", type=kind)
        cell_properties = element("tcPr")
        cell_properties.append(element("shd", val="clear", fill=fill))
        override.append(cell_properties)
        base.element.append(override)
    # Child and parent declarations on the same region cascade by physical edge.
    override = element("tblStylePr", type="firstRow")
    cell_properties = element("tcPr")
    border_properties = element("tcBorders")
    border_properties.append(element("bottom", val="single", sz=12, color="D9D9D9"))
    cell_properties.append(border_properties)
    if args.text:
        existing = next(item for item in child.element if item.tag == qn("w:tblStylePr") and item.get(qn("w:type")) == "firstRow")
        existing.append(cell_properties)
    else:
        override.append(cell_properties)
        child.element.append(override)

table = doc.add_table(rows=5 if args.conditional else 3, cols=3 if args.conditional else 2)
table.style = child
table.autofit = False
# Remove python-docx's direct auto width so inherited physical width is tested.
for item in list(table._tbl.tblPr):
    if item.tag == qn("w:tblW"):
        table._tbl.tblPr.remove(item)
widths = [100, 180, 80] if args.conditional else [120, 240]
for column, width in zip(table.columns, widths):
    column.width = Pt(width)
records = [
    ("Sample", "Observation"),
    ("Sample A", "Measured value retains the inherited fill and comfortable physical cell padding."),
    ("Sample B", "Second editable observation with a longer line that wraps within the inherited table width."),
]
if args.conditional:
    records = [
        ("Sample", "Observation", "Status"),
        ("Sample A", "Editable observation with a direct yellow fill override.", "Open"),
        ("Sample B", "Middle body row retains the even band fill.", "Ready"),
        ("Sample C", "Third body row retains the alternating pale blue band.", "Review"),
        ("Summary", "Final row and corner styles remain distinct.", "Complete"),
    ]
    for item in list(table._tbl.tblPr):
        if item.tag == qn("w:tblLook"):
            table._tbl.tblPr.remove(item)
    table._tbl.tblPr.append(element("tblLook", val="05E0", firstRow=1, lastRow=1, firstColumn=1, lastColumn=1, noHBand=0, noVBand=1))
for row, record in zip(table.rows, records):
    for index, (cell, text) in enumerate(zip(row.cells, record)):
        cell.width = Pt(widths[index])
        cell.text = text
        if args.text:
            cell.paragraphs[0].style = doc.styles["AuditTableBody"]
if not args.text:
    for run in table.cell(0, 0).paragraphs[0].runs + table.cell(0, 1).paragraphs[0].runs:
        run.bold = True
else:
    # A direct black run must beat the base slate colour without disturbing the
    # inherited font or paragraph spacing in this edited yellow cell.
    table.cell(1, 1).paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
# Explicit zero and clear fill are intentional test cases, not missing data.
zero = element("tcMar")
zero.append(element("left", w=0, type="dxa"))
table.cell(1, 0)._tc.get_or_add_tcPr().append(zero)
table.cell(1, 1)._tc.get_or_add_tcPr().append(element("shd", val="clear", fill="FFFFCC"))
table.cell(2, 0)._tc.get_or_add_tcPr().append(element("shd", val="clear", fill="auto"))
doc.add_paragraph("The source uses a named parent and child table style. Export should retain the effective appearance as direct declarations.")
target = Path(args.output)
target.parent.mkdir(parents=True, exist_ok=True)
doc.save(target)
print(target)
