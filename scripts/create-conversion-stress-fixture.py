"""Independent python-docx/Pillow fixture; deliberately never imports Fountain."""
from pathlib import Path
import argparse, hashlib, json, math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from PIL import Image, ImageDraw, ImageFont

parser = argparse.ArgumentParser()
parser.add_argument('--output', required=True)
args = parser.parse_args()
root = Path(args.output); root.mkdir(parents=True, exist_ok=True)
font_path = 'C:/Windows/Fonts/arial.ttf'
def font(size): return ImageFont.truetype(font_path, size)

# Actual numeric plot, with labelled axes and two independently sampled curves.
plot = Image.new('RGB', (1200, 560), 'white'); draw = ImageDraw.Draw(plot)
draw.text((70, 20), 'Cooling curves', fill='black', font=font(34))
for t in range(0, 31, 5):
    x = 110 + t * 31
    draw.line((x, 90, x, 430), fill='#dddddd', width=2)
    draw.text((x-12, 444), str(t), fill='black', font=font(23))
for temp in range(20, 91, 10):
    y = 430 - (temp - 20) * 4.6
    draw.line((110, y, 1040, y), fill='#dddddd', width=2)
    draw.text((55, y-12), str(temp), fill='black', font=font(23))
draw.line((110, 85, 110, 430, 1040, 430), fill='black', width=3)
draw.text((490, 492), 'Time in minutes', fill='black', font=font(26))
draw.text((110, 55), 'Temperature in degrees C', fill='black', font=font(24))
for k, color, name, label_y in [(0.08, '#175e9a', 'Uninsulated', 95), (0.04, '#ad411f', 'Insulated', 140)]:
    points = [(110+t*31, 430-60*math.exp(-k*t)*4.6) for t in range(31)]
    draw.line(points, fill=color, width=5)
    for x,y in points[::5]: draw.ellipse((x-6,y-6,x+6,y+6), fill=color)
    draw.text((800,label_y), name, fill=color, font=font(24))
plot.save(root/'cooling-plot.png')

diagram = Image.new('RGB',(1000,320),'white'); d=ImageDraw.Draw(diagram)
for x, title, fill in [(60,'Uninsulated','#dbe9f5'),(540,'Insulated','#f3e4d6')]:
    d.rounded_rectangle((x,75,x+350,250),radius=18,fill=fill,outline='black',width=4)
    d.text((x+35,270),title,fill='black',font=font(27))
    d.line((x+170,25,x+170,185),fill='#333333',width=8)
    d.ellipse((x+158,173,x+182,197),fill='#ba3024')
    d.text((x+200,22),'Probe',fill='black',font=font(24))
diagram.save(root/'apparatus.jpg',quality=94)
logo=Image.new('RGB',(480,100),'#153b4b'); ld=ImageDraw.Draw(logo)
ld.text((24,26),'THERMAL LAB',fill='white',font=font(38)); logo.save(root/'header-logo.png')

doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.72)
section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Caption'].font.size=Pt(10)
header=section.header.paragraphs[0]
header.add_run().add_picture(str(root/'header-logo.png'),width=Inches(1.25))
header.add_run('   Cooling study 2026')
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Thermal experiments   Page ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text='1'; run.append(text); field.append(run); footer._p.append(field)

doc.add_paragraph('Cooling experiment report','Title')
doc.add_paragraph('Comparison of insulated and uninsulated vessels','Subtitle')
p=doc.add_paragraph('We compare how two vessels cool from 80 degrees C in a room held at 20 degrees C. ')
p.add_run('The insulated vessel cools more slowly.').bold=True
p.add_run(' The chart and measurements below describe the same cooling model.')
doc.add_heading('Experimental setup',1)
doc.add_paragraph('Each vessel contains the same water volume. A probe records temperature every five minutes. Keep the lid position and room conditions unchanged.')
for line in ['Record the initial temperature before starting the timer.','Keep both probes at the same depth.','Read the measurements without moving either vessel.']:
    doc.add_paragraph(line,'List Number')
def picture(path,width,alt):
    shape=doc.add_picture(str(path),width=Inches(width)); shape._inline.docPr.set('descr',alt)
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
picture(root/'apparatus.jpg',5.9,'Two vessels with temperature probes')
doc.add_paragraph('Figure 1  Probe positions in the two vessels','Caption')
picture(root/'cooling-plot.png',6.1,'Cooling plot comparing two temperature series')
doc.add_paragraph('Figure 2  Cooling curves over thirty minutes','Caption')

doc.add_page_break(); doc.add_heading('Model and measurements',1)
doc.add_paragraph('Newton cooling model')
eq1='''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>T(t) = T</m:t></m:r><m:sSub><m:e><m:r><m:t>a</m:t></m:r></m:e><m:sub><m:r><m:t> </m:t></m:r></m:sub></m:sSub><m:r><m:t> + (T</m:t></m:r><m:sSub><m:e><m:r><m:t>0</m:t></m:r></m:e><m:sub><m:r><m:t> </m:t></m:r></m:sub></m:sSub><m:r><m:t> − T_a) </m:t></m:r><m:sSup><m:e><m:r><m:t>e</m:t></m:r></m:e><m:sup><m:r><m:t>−kt</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>'''
# Use genuine subscript structures, not underscore source disguised as equation text.
eq1='''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>T(t) = </m:t></m:r><m:sSub><m:e><m:r><m:t>T</m:t></m:r></m:e><m:sub><m:r><m:t>a</m:t></m:r></m:sub></m:sSub><m:r><m:t> + (</m:t></m:r><m:sSub><m:e><m:r><m:t>T</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub><m:r><m:t> − </m:t></m:r><m:sSub><m:e><m:r><m:t>T</m:t></m:r></m:e><m:sub><m:r><m:t>a</m:t></m:r></m:sub></m:sSub><m:r><m:t>) </m:t></m:r><m:sSup><m:e><m:r><m:t>e</m:t></m:r></m:e><m:sup><m:r><m:t>−kt</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>'''
doc.add_paragraph()._p.append(parse_xml(eq1))
doc.add_paragraph('Half cooling time')
eq2='''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:sSub><m:e><m:r><m:t>t</m:t></m:r></m:e><m:sub><m:r><m:t>1/2</m:t></m:r></m:sub></m:sSub><m:r><m:t> = </m:t></m:r><m:f><m:num><m:r><m:t>ln 2</m:t></m:r></m:num><m:den><m:r><m:t>k</m:t></m:r></m:den></m:f></m:oMath></m:oMathPara>'''
doc.add_paragraph()._p.append(parse_xml(eq2))
doc.add_paragraph('The rate constants are 0.08 and 0.04 per minute. The half cooling times are approximately 8.66 and 17.33 minutes respectively.')
doc.add_heading('Recorded temperatures',1)
table=doc.add_table(rows=2, cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.cell(0,0).merge(table.cell(0,2)).text='Water temperature in degrees C'
for cell,label in zip(table.rows[1].cells,['Time in minutes','Uninsulated','Insulated']): cell.text=label
for t in [0,5,10,15,20,25,30]:
    cells=table.add_row().cells
    for cell,value in zip(cells,[str(t),f'{20+60*math.exp(-.08*t):.1f}',f'{20+60*math.exp(-.04*t):.1f}']): cell.text=value
for i,row in enumerate(table.rows):
    for cell in row.cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=cell._tc.get_or_add_tcPr()
        fill='173B59' if i<2 else ('EFF4F8' if i%2==0 else 'FFFFFF')
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>'))
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            border=OxmlElement('w:'+side); border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'4'); border.set(qn('w:color'),'D9D9D9'); borders.append(border)
        tcPr.append(borders)
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
            for run in p.runs: run.font.color.rgb=RGBColor.from_string('FFFFFF' if i<2 else '000000'); run.font.bold=i<2
doc.add_paragraph('Table 1  Measurements at five minute intervals','Caption')
doc.add_heading('Interpretation',1)
p=doc.add_paragraph('After thirty minutes, the insulated vessel remains approximately 12.6 degrees C warmer. [[FN]]')
p=doc.add_paragraph('Review decision: '); p.add_run('repeat the experiment with a second probe').italic=True
doc.add_paragraph('END OF REPORT')
doc.core_properties.title='Cooling experiment report'; doc.core_properties.author='Conversion audit'
doc.save(root/'cooling-report-base.docx')
manifest={'fixture':'independent-python-docx-v1','expected':{'body_images':2,'header_images':1,'native_equations':2,'tables':1,'footnotes':1,'explicit_page_breaks':1},'assets':{name:hashlib.sha256((root/name).read_bytes()).hexdigest() for name in ['cooling-plot.png','apparatus.jpg','header-logo.png']}}
(root/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
print(root/'cooling-report-base.docx')
